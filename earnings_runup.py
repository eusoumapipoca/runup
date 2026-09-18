#!/usr/bin/env python3
"""
earnings_runup.py — Gerador de candidatos para a Estrategia 4 (run-up pre-resultados)

Fluxo:
  1. Universo base (S&P 500 por defeito, ou ficheiro de tickers)
  2. Calendario de resultados via Finnhub (janela futura configuravel)
  3. Precos via yfinance -> MM20/50/200, ATR(14), volume medio
  4. Gap na reacao ao ultimo anuncio
  5. Datas T-15 / T-12 / SAIDA em sessoes de bolsa (calendario NYSE, feriados incluidos)
  6. CSV pronto para o journal

NAO substitui a verificacao manual da data na pagina de Investor Relations.
O campo 'hour' do Finnhub e frequentemente estimado.

Uso:
  export FINNHUB_TOKEN=xxx
  python earnings_runup.py --weeks 8
  python earnings_runup.py --weeks 8 --universe meus_tickers.txt --out candidatos.csv
"""

import argparse
import os
import sys
import time
from datetime import date, datetime, timedelta

import pandas as pd
import pandas_market_calendars as mcal

# ----------------------------------------------------------------------------
# Parametros da estrategia
# ----------------------------------------------------------------------------
ENTRY_EARLY = 15          # T-15: inicio da janela de entrada
ENTRY_LATE = 12           # T-12: fim da janela de entrada
MIN_AVG_VOLUME = 1_000_000
MIN_PRICE = 10.0
MAX_LAST_GAP_DOWN = -5.0  # gap negativo pior que isto elimina
MAX_HIST_MOVE = 8.0       # movimento medio nos ultimos 4 anuncios; acima elimina
MAX_DRAWDOWN_52S = 15.0   # % abaixo do maximo de 52 semanas; acima elimina.
                          # A MM50/MM200 sao lentas e nao apanham uma correcao
                          # recente com forca (ex: AMAT caiu 20% desde o maximo
                          # de 30-jun e continuava "acima da MM50" nos filtros
                          # antigos). Isto e o que a MM20 apanhava so na saida;
                          # agora tambem filtra a entrada.
ATR_STOP_MULT = 3.0      # stop a 3 ATR: em 15 sessoes, 1.2 ATR era tocado
                         # por ruido em ~55% dos casos; a 3.0 cai para ~20%
MAX_BETA = 1.2           # default; ajustavel via --beta-max (99 desliga o filtro)
SLEEVE_EUR = 2000.0      # capital dedicado a esta estrategia
RISCO_PCT = 0.5          # % do sleeve arriscada por trade
TECTO_POSICAO_PCT = 8.0  # nenhuma posicao isolada acima disto
LIMITE_AGREGADO_PCT = 15.0  # exposicao total simultanea em run-ups
FX_FALLBACK = 1.10       # EUR/USD usado se a cotacao nao for obtida

PRE_JANELA_DIAS = 10     # so consulta historico de anuncios para nomes cuja
                         # janela de entrada abre dentro deste nº de dias

NYSE = mcal.get_calendar("XNYS")


# ----------------------------------------------------------------------------
# Sessoes de bolsa
# ----------------------------------------------------------------------------
def trading_sessions(start: date, end: date) -> list:
    """Lista de sessoes de bolsa NYSE entre duas datas (feriados excluidos)."""
    sched = NYSE.schedule(start_date=start.isoformat(), end_date=end.isoformat())
    return [d.date() for d in sched.index]


def sessions_before(anchor: date, n: int, sessions: list) -> date:
    """Devolve a sessao que esta n sessoes antes de 'anchor'."""
    prior = [s for s in sessions if s < anchor]
    if len(prior) < n:
        return None
    return prior[-n]


def compute_dates(earn_date: date, hour: str, sessions: list) -> dict:
    """
    Calcula janela de entrada e data de saida obrigatoria.

    hour == 'amc'  -> divulga depois do fecho: a ultima sessao para sair e o
                      PROPRIO dia do anuncio.
    hour == 'bmo'  -> divulga antes da abertura: sai no fecho da sessao anterior.
    hour == 'dmh'  -> durante a sessao: trata como bmo (mais conservador).
    """
    h = (hour or "").lower()
    if h == "amc":
        exit_date = earn_date if earn_date in sessions else sessions_before(earn_date, 1, sessions)
        exit_note = "fecho do proprio dia do anuncio (AMC)"
    else:
        exit_date = sessions_before(earn_date, 1, sessions)
        exit_note = "fecho da sessao anterior ao anuncio (BMO/DMH)"

    return {
        "T_minus_15": sessions_before(earn_date, ENTRY_EARLY, sessions),
        "T_minus_12": sessions_before(earn_date, ENTRY_LATE, sessions),
        "data_saida": exit_date,
        "nota_saida": exit_note,
    }


# ----------------------------------------------------------------------------
# Universo
# ----------------------------------------------------------------------------
SP500_CSV = ("https://raw.githubusercontent.com/datasets/s-and-p-500-companies"
             "/master/data/constituents.csv")
UA = {"User-Agent": "Mozilla/5.0 (compatible; earnings-runup/1.0)"}


def load_universe(path: str = None) -> list:
    """
    Universo base. Por ordem de preferencia:
      1. ficheiro local passado em --universe
      2. CSV de constituintes do S&P 500 alojado no GitHub (estavel, sem bloqueios)
      3. Wikipedia, com User-Agent explicito (bloqueia pedidos sem UA -> HTTP 403)
    """
    import io
    import requests

    if path:
        with open(path) as f:
            return sorted({ln.strip().upper() for ln in f if ln.strip()})

    print("A obter constituintes do S&P 500...", file=sys.stderr)
    try:
        r = requests.get(SP500_CSV, headers=UA, timeout=30)
        r.raise_for_status()
        syms = pd.read_csv(io.StringIO(r.text))["Symbol"]
    except Exception as exc:
        print(f"  fonte primaria falhou ({exc}); a tentar Wikipedia...", file=sys.stderr)
        try:
            r = requests.get("https://en.wikipedia.org/wiki/List_of_S%26P_500_companies",
                             headers=UA, timeout=30)
            r.raise_for_status()
            syms = pd.read_html(io.StringIO(r.text))[0]["Symbol"]
        except Exception as exc2:
            sys.exit(f"Nao foi possivel obter o universo ({exc2}).\n"
                     f"Alternativa: cria um ficheiro de texto com um ticker por linha "
                     f"e corre com --universe tickers.txt")

    tickers = syms.astype(str).str.strip().str.replace(".", "-", regex=False)
    return sorted(set(tickers))


# ----------------------------------------------------------------------------
# Calendario Finnhub
# ----------------------------------------------------------------------------
def fetch_calendar(token: str, start: date, end: date, chunk_days: int = 7) -> pd.DataFrame:
    """
    Calendario de resultados, pedido em fatias curtas.

    O tier gratuito do Finnhub trunca respostas grandes sem avisar: um pedido
    de 8 semanas devolve MENOS linhas do que um de 2 semanas. Fatiar em blocos
    de ~7 dias mantem cada resposta abaixo do tecto.
    """
    import requests

    frames, cur = [], start
    while cur <= end:
        stop = min(cur + timedelta(days=chunk_days - 1), end)
        try:
            r = requests.get(
                "https://finnhub.io/api/v1/calendar/earnings",
                params={"from": cur.isoformat(), "to": stop.isoformat(), "token": token},
                timeout=30,
            )
            r.raise_for_status()
            rows = r.json().get("earningsCalendar", []) or []
            if rows:
                frames.append(pd.DataFrame(rows))
        except Exception as exc:
            print(f"  fatia {cur}..{stop} falhou: {exc}", file=sys.stderr)
        cur = stop + timedelta(days=1)
        time.sleep(1.1)          # tier gratuito: 60 pedidos/minuto

    if not frames:
        return pd.DataFrame()
    df = pd.concat(frames, ignore_index=True)
    df["date"] = pd.to_datetime(df["date"]).dt.date
    return df.drop_duplicates(subset=["symbol", "date"])


# ----------------------------------------------------------------------------
# Precos e indicadores
# ----------------------------------------------------------------------------
def fetch_prices(tickers: list) -> dict:
    import yfinance as yf

    out = {}
    batch = 60
    for i in range(0, len(tickers), batch):
        chunk = tickers[i:i + batch]
        try:
            data = yf.download(chunk, period="15mo", interval="1d",
                               group_by="ticker", auto_adjust=False,
                               progress=False, threads=True)
        except Exception as exc:
            print(f"  lote {i//batch + 1} falhou: {exc}", file=sys.stderr)
            continue

        if data is None or data.empty:
            continue

        for t in chunk:
            try:
                # com 1 ticker o yfinance devolve colunas simples;
                # com varios devolve MultiIndex (ticker, campo)
                if isinstance(data.columns, pd.MultiIndex):
                    if t not in data.columns.get_level_values(0):
                        continue
                    df = data[t].dropna()
                else:
                    df = data.dropna()
                if len(df) > 200 and {"Open", "High", "Low", "Close", "Volume"} <= set(df.columns):
                    out[t] = df
            except Exception:
                continue
        time.sleep(1)
    return out


def atr(df: pd.DataFrame, n: int = 14) -> float:
    h, l, c = df["High"], df["Low"], df["Close"]
    pc = c.shift(1)
    tr = pd.concat([h - l, (h - pc).abs(), (l - pc).abs()], axis=1).max(axis=1)
    return float(tr.rolling(n).mean().iloc[-1])


def last_gap_pct(df: pd.DataFrame, last_earn: date, hour: str) -> float:
    """
    Gap % na primeira sessao a reagir ao ultimo anuncio.

    Se 'hour' for desconhecido (None), nao sabemos se a reacao foi no proprio
    dia (BMO) ou no seguinte (AMC). Nesse caso mede os dois e devolve o de
    maior movimento absoluto — na pratica, a sessao da reacao.
    """
    if last_earn is None:
        return None
    idx = [d.date() for d in df.index]
    try:
        pos = next(i for i, d in enumerate(idx) if d >= last_earn)
    except StopIteration:
        return None

    def gap_em(i):
        if i < 1 or i >= len(df):
            return None
        prev_close = float(df["Close"].iloc[i - 1])
        open_px = float(df["Open"].iloc[i])
        return round((open_px / prev_close - 1) * 100, 2)

    h = (hour or "").lower()
    if h == "amc":
        return gap_em(pos + 1)
    if h in ("bmo", "dmh"):
        return gap_em(pos)

    cands = [g for g in (gap_em(pos), gap_em(pos + 1)) if g is not None]
    return max(cands, key=abs) if cands else None


def past_earnings_via_yf(ticker: str, n: int = 5) -> list:
    """
    Ate n datas de resultados ja divulgadas, da mais recente para tras.
    Fonte independente do Finnhub, cujo tier gratuito trunca o historico.
    """
    import yfinance as yf

    try:
        df = yf.Ticker(ticker).get_earnings_dates(limit=16)
        if df is None or df.empty:
            return []
        idx = pd.to_datetime(df.index).tz_localize(None)
        passadas = sorted({d.date() for d in idx if d.date() < date.today()})
        return passadas[-n:][::-1]
    except Exception:
        return []


def move_on_earnings(df: pd.DataFrame, earn: date, hour: str = None) -> float:
    """
    Movimento total (fecho vs fecho anterior) na sessao que reagiu ao anuncio.
    Diferente de last_gap_pct, que mede so a abertura: aqui interessa a
    amplitude completa da reacao, que e o que o movimento implicito estima.
    """
    if earn is None:
        return None
    idx = [d.date() for d in df.index]
    try:
        pos = next(i for i, d in enumerate(idx) if d >= earn)
    except StopIteration:
        return None

    def mv(i):
        if i < 1 or i >= len(df):
            return None
        prev_close = float(df["Close"].iloc[i - 1])
        close = float(df["Close"].iloc[i])
        return round((close / prev_close - 1) * 100, 2)

    h = (hour or "").lower()
    if h == "amc":
        return mv(pos + 1)
    if h in ("bmo", "dmh"):
        return mv(pos)
    cands = [m for m in (mv(pos), mv(pos + 1)) if m is not None]
    return max(cands, key=abs) if cands else None


def historic_move(df: pd.DataFrame, datas: list, n: int = 4):
    """
    Media dos movimentos absolutos nos ultimos n anuncios.
    Substitui a leitura manual do movimento implicito nas opcoes.
    Devolve (media, lista_dos_movimentos).
    """
    movs = []
    for d in datas[:n]:
        m = move_on_earnings(df, d)
        if m is not None:
            movs.append(m)
    if not movs:
        return None, []
    return round(sum(abs(m) for m in movs) / len(movs), 2), movs


def fetch_benchmark(period: str = "15mo") -> pd.Series:
    """Retornos diarios do SPY, para calculo de beta."""
    import yfinance as yf

    try:
        d = yf.download("SPY", period=period, interval="1d",
                        auto_adjust=False, progress=False)
        if d is None or d.empty:
            return None
        c = d["Close"]
        if isinstance(c, pd.DataFrame):
            c = c.iloc[:, 0]
        return c.pct_change().dropna()
    except Exception as exc:
        print(f"  benchmark indisponivel ({exc}); beta nao sera calculado",
              file=sys.stderr)
        return None


def beta_vs(df: pd.DataFrame, bench: pd.Series, dias: int = 252) -> float:
    """
    Beta sobre os ultimos ~12 meses de retornos diarios.
    Importa nesta estrategia porque a posicao fica 15 sessoes no mercado:
    com beta alto, o movimento do indice afoga o sinal de run-up.
    """
    if bench is None:
        return None
    try:
        r = df["Close"].pct_change().dropna()
        j = pd.concat([r, bench], axis=1, join="inner").dropna()
        j = j.iloc[-dias:]
        if len(j) < 60:
            return None
        x = j.iloc[:, 1]
        y = j.iloc[:, 0]
        var = float(x.var())
        if var == 0:
            return None
        return round(float(y.cov(x)) / var, 2)
    except Exception:
        return None


def technicals(df: pd.DataFrame) -> dict:
    c = df["Close"]
    price = float(c.iloc[-1])
    sma20 = float(c.rolling(20).mean().iloc[-1])
    sma50 = float(c.rolling(50).mean().iloc[-1])
    sma200 = float(c.rolling(200).mean().iloc[-1])
    a = atr(df)

    # Maximo dos ultimos ~252 sessoes (52 semanas), ou o historico disponivel
    # se for mais curto.
    janela_52s = c.tail(252)
    maximo_52s = float(janela_52s.max())
    drawdown_52s = round((price / maximo_52s - 1) * 100, 2) if maximo_52s else None

    return {
        "preco": round(price, 2),
        "mm20": round(sma20, 2),
        "mm50": round(sma50, 2),
        "mm200": round(sma200, 2),
        "atr14": round(a, 2),
        "vol_medio_50d": int(df["Volume"].rolling(50).mean().iloc[-1]),
        "stop": round(price - ATR_STOP_MULT * a, 2),
        # Preco > MM20 entrou na entrada, nao so na saida: a MM20 reage a uma
        # correcao recente muito mais depressa que a MM50/MM200, que podem
        # continuar "otimistas" semanas depois de a tendencia ja ter virado.
        "tendencia_ok": bool(price > sma20 and price > sma50 and sma50 > sma200),
        "acima_mm20": bool(price > sma20),
        "maximo_52s": round(maximo_52s, 2),
        "drawdown_52s_%": drawdown_52s,
    }




def obter_fx_eurusd(fallback: float = FX_FALLBACK) -> tuple:
    """
    Cotacao EUR/USD. Devolve (taxa, origem).
    Necessaria porque os precos vem em USD e o sleeve esta em EUR.
    """
    import yfinance as yf
    try:
        d = yf.download("EURUSD=X", period="5d", interval="1d",
                        auto_adjust=False, progress=False)
        if d is not None and not d.empty:
            c = d["Close"]
            if isinstance(c, pd.DataFrame):
                c = c.iloc[:, 0]
            taxa = float(c.dropna().iloc[-1])
            if 0.5 < taxa < 2.0:
                return round(taxa, 4), "mercado"
    except Exception:
        pass
    return fallback, "fallback"


def dimensionar(preco_usd, stop_usd, sleeve_eur, risco_pct, tecto_pct, fx):
    """
    Traduz a distancia ao stop num tamanho de posicao concreto.

    O principio: o risco em euros e fixo (uma % do sleeve), e o tamanho da
    posicao ajusta-se a distancia do stop. Stop largo -> posicao pequena.
    Devolve None se faltarem dados.
    """
    if not preco_usd or not stop_usd or preco_usd <= stop_usd:
        return None

    dist_pct = (preco_usd - stop_usd) / preco_usd
    risco_eur = sleeve_eur * risco_pct / 100.0
    posicao_eur = risco_eur / dist_pct

    tecto_eur = sleeve_eur * tecto_pct / 100.0
    limitado = posicao_eur > tecto_eur
    if limitado:
        posicao_eur = tecto_eur

    preco_eur = preco_usd / fx
    n_acoes = posicao_eur / preco_eur

    return {
        "distancia_stop_pct": round(dist_pct * 100, 2),
        "risco_alvo_eur": round(risco_eur, 2),
        "posicao_eur": round(posicao_eur, 2),
        "n_acoes": round(n_acoes, 4),
        "preco_eur": round(preco_eur, 2),
        "pct_do_sleeve": round(posicao_eur / sleeve_eur * 100, 2),
        # Quando o tecto morde, o risco efetivo fica ABAIXO do alvo -- e o
        # comportamento correto: o tecto e uma restricao adicional, nao um alvo.
        "risco_efetivo_eur": round(posicao_eur * dist_pct, 2),
        "limitado_pelo_tecto": limitado,
    }



# ----------------------------------------------------------------------------
# Historico — le o journal.xlsx e resume o desempenho
# ----------------------------------------------------------------------------
def ler_journal(caminho: str, fx: float) -> dict:
    """
    Le o journal.xlsx e devolve um resumo do desempenho por versao do sistema.

    Le apenas as colunas de INPUT e recalcula tudo em Python, em vez de confiar
    nos valores em cache das formulas do Excel -- assim funciona mesmo que o
    ficheiro nunca tenha sido aberto no Excel depois da ultima edicao.
    """
    try:
        from openpyxl import load_workbook
    except ImportError:
        return None

    if not os.path.exists(caminho):
        return None

    try:
        wb = load_workbook(caminho, data_only=True)
        if "Journal" not in wb.sheetnames:
            return None
        ws = wb["Journal"]
        hdr = [c.value for c in ws[1]]
        idx = {h: i for i, h in enumerate(hdr) if h}

        def val(row, nome):
            i = idx.get(nome)
            return row[i] if i is not None and i < len(row) else None

        def como_data(v):
            """O Excel devolve datas como datetime, mas se a celula foi
            escrita como texto vem string. Aceita ambos."""
            if isinstance(v, datetime):
                return v.date()
            if isinstance(v, date):
                return v
            if isinstance(v, str) and v.strip():
                try:
                    return datetime.fromisoformat(v.strip()[:10]).date()
                except ValueError:
                    return None
            return None

        trades = []
        for row in ws.iter_rows(min_row=2, values_only=True):
            tk = val(row, "ticker")
            if not tk:
                continue

            entrada = val(row, "preco_entrada_usd")
            saida = val(row, "preco_saida_usd")
            tam = val(row, "tamanho_acoes")
            d_ent = como_data(val(row, "data_entrada"))
            d_sai = como_data(val(row, "data_saida_real"))
            spy = val(row, "retorno_spy_%")

            fechado = bool(saida and entrada)
            investido = (tam * entrada / fx) if (tam and entrada) else None
            ret_pct = ((saida / entrada) - 1) if fechado else None
            ret_eur = (investido * ret_pct) if (investido and ret_pct is not None) else None
            alfa = (ret_pct - spy) if (ret_pct is not None and spy is not None) else None

            dias = (d_sai - d_ent).days if (d_ent and d_sai) else None

            trades.append({
                "ticker": tk,
                "versao_sistema": val(row, "versao_sistema"),
                "setor": val(row, "setor"),
                "beta": val(row, "beta"),
                "data_entrada": d_ent.isoformat() if d_ent else None,
                "data_saida_real": d_sai.isoformat() if d_sai else None,
                "data_saida_planeada": (lambda x: x.isoformat() if x else None)(
                    como_data(val(row, "data_saida_planeada"))),
                "preco_entrada_usd": entrada,
                "preco_saida_usd": saida,
                "tamanho_acoes": tam,
                "valor_investido_eur": round(investido, 2) if investido else None,
                "motivo_saida": val(row, "motivo_saida"),
                "dias_em_posicao": dias,
                "retorno_pct": round(ret_pct * 100, 2) if ret_pct is not None else None,
                "retorno_eur": round(ret_eur, 2) if ret_eur is not None else None,
                "retorno_spy_pct": round(spy * 100, 2) if spy is not None else None,
                "alfa_pct": round(alfa * 100, 2) if alfa is not None else None,
                "fechado": fechado,
                "nota": val(row, "nota"),
            })
    except Exception as exc:
        print(f"  journal nao lido ({exc})", file=sys.stderr)
        return None

    def resumir(lista):
        fech = [t for t in lista if t["fechado"] and t["retorno_pct"] is not None]
        if not fech:
            return {
                "n_fechados": 0, "n_abertos": len([t for t in lista if not t["fechado"]]),
                "taxa_acerto_pct": None, "retorno_medio_pct": None,
                "ganho_medio_pct": None, "perda_media_pct": None,
                "racio_ganho_perda": None, "resultado_total_eur": 0.0,
                "alfa_medio_pct": None, "dias_medios": None,
                "amostra_suficiente": False, "faltam_para_30": 30,
                "veredicto": "aguarda amostra", "por_motivo": {},
            }
        rets = [t["retorno_pct"] for t in fech]
        ganhos = [r for r in rets if r > 0]
        perdas = [r for r in rets if r < 0]
        eur = [t["retorno_eur"] for t in fech if t["retorno_eur"] is not None]
        alfas = [t["alfa_pct"] for t in fech if t["alfa_pct"] is not None]
        dias = [t["dias_em_posicao"] for t in fech if t["dias_em_posicao"] is not None]

        g_med = sum(ganhos) / len(ganhos) if ganhos else None
        p_med = sum(perdas) / len(perdas) if perdas else None
        racio = round(abs(g_med / p_med), 2) if (g_med and p_med) else None
        acerto = round(len(ganhos) / len(fech) * 100, 1)
        alfa_m = round(sum(alfas) / len(alfas), 2) if alfas else None

        motivos = {}
        for t in fech:
            m = t.get("motivo_saida") or "sem_motivo"
            motivos[m] = motivos.get(m, 0) + 1

        # Criterios de falsificacao definidos antes de haver dados
        if len(fech) < 30:
            veredicto = "aguarda amostra"
        elif acerto < 48 or (racio is not None and racio < 1.3) or \
             (alfa_m is not None and alfa_m <= 0):
            veredicto = "FALHA"
        else:
            veredicto = "PASSA"

        return {
            "n_fechados": len(fech),
            "n_abertos": len([t for t in lista if not t["fechado"]]),
            "taxa_acerto_pct": acerto,
            "retorno_medio_pct": round(sum(rets) / len(rets), 2),
            "ganho_medio_pct": round(g_med, 2) if g_med else None,
            "perda_media_pct": round(p_med, 2) if p_med else None,
            "racio_ganho_perda": racio,
            "resultado_total_eur": round(sum(eur), 2) if eur else 0.0,
            "alfa_medio_pct": alfa_m,
            "dias_medios": round(sum(dias) / len(dias), 1) if dias else None,
            "amostra_suficiente": len(fech) >= 30,
            "faltam_para_30": max(0, 30 - len(fech)),
            "veredicto": veredicto,
            "por_motivo": motivos,
        }

    versoes = sorted({t["versao_sistema"] for t in trades if t["versao_sistema"]})
    return {
        "trades": trades,
        "resumo_por_versao": {v: resumir([t for t in trades if t["versao_sistema"] == v])
                              for v in versoes},
        "resumo_todas": resumir(trades),
        "versao_atual": versoes[-1] if versoes else None,
    }


# ----------------------------------------------------------------------------
# JSON para o dashboard
# ----------------------------------------------------------------------------
def _limpa(v):
    """Converte tipos numpy/pandas/date para algo serializavel em JSON."""
    import numpy as np
    if v is None:
        return None
    if isinstance(v, (np.bool_, bool)):
        return bool(v)
    if isinstance(v, (np.integer,)):
        return int(v)
    if isinstance(v, (np.floating, float)):
        return None if pd.isna(v) else round(float(v), 4)
    if isinstance(v, (date, datetime)):
        return v.isoformat()
    if isinstance(v, pd.Timestamp):
        return v.date().isoformat()
    if pd.isna(v):
        return None
    return str(v)


def escrever_json(args, out, today, funil, ultima_data, atraso, fx, fx_origem,
                  historico=None):
    """
    Escreve o JSON que alimenta o dashboard.

    Estrutura:
      meta       - quando correu, com que parametros, frescura dos dados
      funil      - quantos sobreviveram a cada fase (universo -> aprovados)
      regras     - os limiares em vigor, para o dashboard os poder mostrar
      candidatos - lista completa, cada um com todos os campos calculados
    """
    import json

    registos = []
    for _, r in out.iterrows():
        d = {k: _limpa(v) for k, v in r.to_dict().items()}
        # lista de movimentos como array em vez de string
        movs = d.get("movs_4_ultimos") or ""
        d["movs_4_ultimos"] = [float(x) for x in movs.split(",") if x] if movs else []
        # dimensionamento: da distancia ao stop para euros e nº de acoes
        dim = dimensionar(d.get("preco"), d.get("stop"), args.sleeve,
                          args.risco, args.tecto, fx)
        d["dimensionamento"] = dim
        if dim:
            d["distancia_stop_pct"] = dim["distancia_stop_pct"]
        # motivos como array
        mot = d.get("motivo") or ""
        d["motivo"] = [m for m in mot.split(",") if m]
        registos.append(d)

    # Se entrasse em todos os aprovados que estao em janela hoje, quanto
    # ficaria exposto? E o numero que diz se o limite agregado e respeitado.
    em_janela = [c for c in registos
                 if c.get("janela_aberta_hoje") and c.get("estado") == "OK"
                 and c.get("dimensionamento")]
    total_eur = round(sum(c["dimensionamento"]["posicao_eur"] for c in em_janela), 2)
    limite_eur = round(args.sleeve * args.limite_agregado / 100, 2)
    agregado = {
        "n_posicoes": len(em_janela),
        "exposicao_total_eur": total_eur,
        "pct_do_sleeve": round(total_eur / args.sleeve * 100, 2) if args.sleeve else None,
        "limite_eur": limite_eur,
        "excede_limite": total_eur > limite_eur,
        "tickers": [c["ticker"] for c in em_janela],
    }

    payload = {
        "meta": {
            "gerado_em": datetime.now().isoformat(timespec="seconds"),
            "data_corrida": today.isoformat(),
            "semanas_analisadas": args.weeks,
            "ultima_data_preco": ultima_data.isoformat() if ultima_data else None,
            "sessoes_de_atraso": atraso,
            "dados_frescos": (atraso == 0) if atraso is not None else None,
        },
        "regras": {
            "entrada_sessoes_antes": [ENTRY_EARLY, ENTRY_LATE],
            "beta_max": args.beta_max,
            "atr_mult_stop": args.atr_mult,
            "max_drawdown_52s_pct": MAX_DRAWDOWN_52S,
            "max_movimento_historico_pct": MAX_HIST_MOVE,
            "max_gap_negativo_pct": MAX_LAST_GAP_DOWN,
            "min_volume_medio": MIN_AVG_VOLUME,
            "min_preco": MIN_PRICE,
        },
        "capital": {
            "sleeve_eur": args.sleeve,
            "risco_por_trade_pct": args.risco,
            "risco_por_trade_eur": round(args.sleeve * args.risco / 100, 2),
            "tecto_posicao_pct": args.tecto,
            "tecto_posicao_eur": round(args.sleeve * args.tecto / 100, 2),
            "limite_agregado_pct": args.limite_agregado,
            "limite_agregado_eur": round(args.sleeve * args.limite_agregado / 100, 2),
            "fx_eurusd": fx,
            "fx_origem": fx_origem,
        },
        "agregado_se_entrar_em_todos": agregado,
        "historico": historico,
        "funil": funil,
        "candidatos": registos,
    }

    with open(args.json, "w", encoding="utf-8") as f:
        json.dump(payload, f, ensure_ascii=False, indent=2)

    n_jan = sum(1 for c in registos if c.get("janela_aberta_hoje"))
    print(f"{len(registos)} candidatos ({n_jan} em janela) -> {args.json}",
          file=sys.stderr)
    return payload



# ----------------------------------------------------------------------------
# Dashboard HTML — ficheiro unico, dados embutidos, sem dependencias
# ----------------------------------------------------------------------------
CAMPOS_MONETARIOS_TRADE = ("valor_investido_eur", "retorno_eur", "tamanho_acoes",
                           "preco_entrada_usd", "preco_saida_usd")
CAMPOS_MONETARIOS_DIM = ("posicao_eur", "n_acoes", "risco_alvo_eur",
                         "risco_efetivo_eur", "preco_eur")


def _versao_publica(payload: dict) -> dict:
    """
    Copia do payload sem valores monetarios absolutos.

    Percentagens ficam: dizem se o metodo funciona. Euros e nº de acoes saem:
    revelariam o tamanho do capital, por deducao mesmo quando nao declarado.
    """
    import copy
    p = copy.deepcopy(payload)

    cap = p.get("capital") or {}
    for k in ("sleeve_eur", "risco_por_trade_eur", "tecto_posicao_eur",
              "limite_agregado_eur"):
        cap.pop(k, None)

    ag = p.get("agregado_se_entrar_em_todos") or {}
    for k in ("exposicao_total_eur", "limite_eur"):
        ag.pop(k, None)

    for c in p.get("candidatos", []):
        dim = c.get("dimensionamento")
        if isinstance(dim, dict):
            for k in CAMPOS_MONETARIOS_DIM:
                dim.pop(k, None)

    h = p.get("historico")
    if isinstance(h, dict):
        for t in h.get("trades", []):
            for k in CAMPOS_MONETARIOS_TRADE:
                t.pop(k, None)
        for bloco in list(h.get("resumo_por_versao", {}).values()) + [h.get("resumo_todas")]:
            if isinstance(bloco, dict):
                bloco.pop("resultado_total_eur", None)
    return p


def escrever_html(payload: dict, caminho: str, template: str, publico: bool = True):
    """Injeta os dados no template e grava o HTML."""
    import json

    dados = _versao_publica(payload) if publico else payload
    html = template.replace("__DADOS__", json.dumps(dados, ensure_ascii=False))
    with open(caminho, "w", encoding="utf-8") as f:
        f.write(html)
    print(f"Dashboard {'publico' if publico else 'completo'} -> {caminho}",
          file=sys.stderr)



def exportar_historico(historico: dict, caminho: str):
    """
    Escreve o historico do journal num JSON ja SEM valores monetarios.

    Existe para o repositorio publico: o journal.xlsx tem euros e nao pode la
    estar, mas o GitHub Actions precisa de alguma fonte para a seccao de
    historico. Este ficheiro e seguro para um repo publico.
    """
    import json

    if not historico:
        print("Sem historico para exportar", file=sys.stderr)
        return

    limpo = _versao_publica({"historico": historico})["historico"]
    with open(caminho, "w", encoding="utf-8") as f:
        json.dump(limpo, f, ensure_ascii=False, indent=2)
    n = len(limpo.get("trades", []))
    print(f"Historico ({n} trades, sem valores monetarios) -> {caminho}",
          file=sys.stderr)


def carregar_historico(args, fx):
    """
    Historico para o dashboard. Prefere o journal.xlsx (local, completo); se
    nao existir, usa o historico.json exportado (repo publico / CI).
    """
    import json

    if os.path.exists(args.journal):
        return ler_journal(args.journal, fx), "journal"
    if os.path.exists(args.historico):
        try:
            with open(args.historico, encoding="utf-8") as f:
                return json.load(f), "historico.json"
        except Exception as exc:
            print(f"  historico.json ilegivel ({exc})", file=sys.stderr)
    return None, None


# ----------------------------------------------------------------------------
# Main
# ----------------------------------------------------------------------------
def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--weeks", type=int, default=8, help="janela futura em semanas")
    ap.add_argument("--universe", help="ficheiro com um ticker por linha")
    ap.add_argument("--out", default="candidatos.csv")
    ap.add_argument("--token", default=os.getenv("FINNHUB_TOKEN"))
    ap.add_argument("--beta-max", type=float, default=MAX_BETA,
                    help="beta maximo; usa 99 para desligar o filtro (default: %(default)s)")
    ap.add_argument("--atr-mult", type=float, default=ATR_STOP_MULT,
                    help="multiplicador do ATR para o stop (default: %(default)s)")
    ap.add_argument("--json", default="dashboard.json",
                    help="ficheiro JSON para o dashboard (default: %(default)s)")
    ap.add_argument("--sleeve", type=float, default=SLEEVE_EUR,
                    help="capital dedicado a esta estrategia, em EUR (default: %(default)s)")
    ap.add_argument("--risco", type=float, default=RISCO_PCT,
                    help="%% do sleeve arriscada por trade (default: %(default)s)")
    ap.add_argument("--tecto", type=float, default=TECTO_POSICAO_PCT,
                    help="%% maxima do sleeve numa unica posicao (default: %(default)s)")
    ap.add_argument("--limite-agregado", type=float, default=LIMITE_AGREGADO_PCT,
                    help="%% maxima do sleeve em posicoes simultaneas (default: %(default)s)")
    ap.add_argument("--fx", type=float, default=None,
                    help="taxa EUR/USD; se omitida, e obtida automaticamente")
    ap.add_argument("--journal", default="journal.xlsx",
                    help="journal.xlsx a resumir no dashboard (default: %(default)s)")
    ap.add_argument("--html", default="index.html",
                    help="dashboard HTML publico (default: %(default)s)")
    ap.add_argument("--template", default="template.html",
                    help="template do dashboard (default: %(default)s)")
    ap.add_argument("--historico", default="historico.json",
                    help="JSON do historico sem valores monetarios, para o repo publico")
    ap.add_argument("--exportar-historico", action="store_true",
                    help="escreve o historico.json a partir do journal e sai")
    ap.add_argument("--all", action="store_true",
                    help="inclui tambem os que falham filtros, com o motivo")
    args = ap.parse_args()

    # exportar o historico nao precisa de rede nem de token
    if args.exportar_historico:
        h = ler_journal(args.journal, args.fx or FX_FALLBACK)
        if not h:
            sys.exit(f"Nao foi possivel ler {args.journal}")
        exportar_historico(h, args.historico)
        return

    if not args.token:
        sys.exit("Falta o token. Define FINNHUB_TOKEN ou usa --token")

    today = date.today()
    end = today + timedelta(weeks=args.weeks)
    sessions = trading_sessions(today - timedelta(days=500), end + timedelta(days=30))

    universe = load_universe(args.universe)
    print(f"Universo: {len(universe)} tickers", file=sys.stderr)

    cal = fetch_calendar(args.token, today, end)
    if cal.empty:
        sys.exit("Calendario vazio.")
    cal = cal[cal["symbol"].isin(universe)].copy()
    cal = cal.sort_values("date").drop_duplicates("symbol", keep="first")
    print(f"Com anuncio na janela: {len(cal)}", file=sys.stderr)

    past = fetch_calendar(args.token, today - timedelta(days=200), today - timedelta(days=1))
    last_map = {}
    if not past.empty:
        past = past.sort_values("date")
        for sym, g in past.groupby("symbol"):
            last_map[sym] = (g["date"].iloc[-1], g["hour"].iloc[-1])
    print(f"Historico de anuncios via Finnhub: {len(last_map)} simbolos", file=sys.stderr)
    if not last_map:
        print("  (vazio — o gap sera obtido via yfinance para os sobreviventes)",
              file=sys.stderr)

    tickers = cal["symbol"].tolist()
    prices = fetch_prices(tickers)
    print(f"Com historico de precos: {len(prices)}", file=sys.stderr)
    ultima, atraso = None, None
    if prices:
        ultima = max(df.index[-1].date() for df in prices.values())
        atraso = len([d for d in sessions if ultima < d <= today])
        aviso = "" if atraso == 0 else f"  <-- {atraso} sessao(oes) de atraso!"
        print(f"Ultimo fecho nos dados: {ultima}{aviso}", file=sys.stderr)

    fx, fx_origem = (args.fx, "manual") if args.fx else obter_fx_eurusd()
    print(f"EUR/USD: {fx} ({fx_origem})  |  sleeve {args.sleeve:.0f} EUR, "
          f"risco {args.risco}% = {args.sleeve*args.risco/100:.2f} EUR/trade",
          file=sys.stderr)

    bench = fetch_benchmark()
    if bench is not None:
        print(f"Benchmark SPY: {len(bench)} retornos diarios", file=sys.stderr)

    rows = []
    n_lento = 0
    n_perto = sum(
        1 for _, rr in cal.iterrows()
        if (lambda dd: dd["T_minus_15"] is not None
                       and dd["T_minus_12"] is not None
                       and dd["T_minus_12"] >= today
                       and (dd["T_minus_15"] - today).days <= PRE_JANELA_DIAS)
           (compute_dates(rr["date"], rr.get("hour"), sessions))
    )
    print(f"Perto da janela (consulta lenta): {n_perto}", file=sys.stderr)
    for _, r in cal.iterrows():
        sym = r["symbol"]
        if sym not in prices:
            continue
        df = prices[sym]
        d = compute_dates(r["date"], r.get("hour"), sessions)
        t = technicals(df)
        t["stop"] = round(t["preco"] - args.atr_mult * t["atr14"], 2)

        # --- filtros baratos primeiro -------------------------------------
        motivos = []
        if not t["tendencia_ok"]:
            motivos.append("tendencia")
        if t["vol_medio_50d"] < MIN_AVG_VOLUME:
            motivos.append("liquidez")
        if t["preco"] < MIN_PRICE:
            motivos.append("preco")
        dd = t.get("drawdown_52s_%")
        if dd is not None and dd < -MAX_DRAWDOWN_52S:
            motivos.append("drawdown_52s")
        if d["T_minus_15"] is None or d["data_saida"] is None:
            motivos.append("datas")

        b = beta_vs(df, bench)
        if b is not None and b > args.beta_max:
            motivos.append("beta")

        # --- gap e movimento historico ------------------------------------
        # Consulta lenta (1 chamada yfinance por ticker). So vale a pena para
        # quem esta na janela ou perto dela: um nome que reporta daqui a 6
        # semanas volta a ser avaliado numa corrida futura.
        gap = None
        mov_hist = None
        movs = []
        reacao = "nao_aplicavel"

        perto = (
            d["T_minus_15"] is not None
            and d["T_minus_12"] is not None
            and d["T_minus_12"] >= today                       # janela ainda nao fechou
            and (d["T_minus_15"] - today).days <= PRE_JANELA_DIAS  # e abre em breve
        )

        if not motivos and not perto:
            reacao = "pendente"

        if not motivos and perto:
            n_lento += 1
            print(f"  [{n_lento}/{n_perto}] historico de {sym}...".ljust(58),
                  end="\r", file=sys.stderr, flush=True)
            datas = past_earnings_via_yf(sym)
            le, lh = last_map.get(sym, (None, None))
            if le is None and datas:
                le, lh = datas[0], None
            gap = last_gap_pct(df, le, lh)
            if gap is not None and gap < MAX_LAST_GAP_DOWN:
                motivos.append("gap_anterior")

            mov_hist, movs = historic_move(df, datas)
            if mov_hist is not None and mov_hist > MAX_HIST_MOVE:
                motivos.append("mov_historico")

            reacao = "aplicados" if (gap is not None or mov_hist is not None) else "sem_dados"

        na_janela = (
            d["T_minus_15"] is not None
            and d["T_minus_15"] <= today <= d["T_minus_12"]
        )

        rows.append({
            "ticker": sym,
            "data_anuncio": r["date"],
            "hora": r.get("hour"),
            "T_minus_15": d["T_minus_15"],
            "T_minus_12": d["T_minus_12"],
            "data_saida": d["data_saida"],
            "nota_saida": d["nota_saida"],
            "janela_aberta_hoje": na_janela,
            **t,
            "data_preco": df.index[-1].date(),
            "beta": b,
            "filtros_reacao": reacao,
            "gap_ultimo_anuncio_%": gap,
            "mov_historico_%": mov_hist,
            "movs_4_ultimos": ",".join(str(m) for m in movs) if movs else "",
            "eps_estimate": r.get("epsEstimate"),
            "estado": "OK" if not motivos else "FALHA",
            "motivo": ",".join(motivos),
            "verificar_IR": "PENDENTE",
            "zacks_rank": "",
        })

    print(" " * 60, end="\r", file=sys.stderr)
    print(f"Historico consultado para {n_lento} candidatos".ljust(58), file=sys.stderr)

    out = pd.DataFrame(rows).sort_values(["data_anuncio", "ticker"])
    if not args.all:
        out = out[out["estado"] == "OK"]
    out.to_csv(args.out, index=False)

    # ---- JSON para o dashboard -------------------------------------------
    historico, origem = carregar_historico(args, fx)
    if historico:
        rt = historico["resumo_todas"]
        print(f"Journal: {rt['n_fechados']} fechados, {rt['n_abertos']} abertos "
              f"(fonte: {origem})", file=sys.stderr)

    payload = escrever_json(args, out, today, {
        "universo": len(universe),
        "com_anuncio": len(cal),
        "com_precos": len(prices),
        "perto_da_janela": n_perto,
        "historico_consultado": n_lento,
        "aprovados": int((out["estado"] == "OK").sum()) if len(out) else 0,
    }, ultima, atraso, fx, fx_origem, historico)

    if os.path.exists(args.template):
        with open(args.template, encoding="utf-8") as f:
            escrever_html(payload, args.html, f.read(), publico=True)
    else:
        print(f"Template {args.template} nao encontrado — HTML nao gerado",
              file=sys.stderr)

    print(f"\n{len(out)} linhas -> {args.out}", file=sys.stderr)
    abertos = out[out["janela_aberta_hoje"]] if "janela_aberta_hoje" in out else pd.DataFrame()
    print(f"Em janela de entrada HOJE: {len(abertos)}", file=sys.stderr)
    if len(abertos):
        print(abertos[["ticker", "data_anuncio", "hora", "data_saida",
                       "preco", "stop"]].to_string(index=False), file=sys.stderr)


if __name__ == "__main__":
    main()
